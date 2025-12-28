import os
from typing import Optional
from pypdf import PdfReader
import docx
import openpyxl

def extract_text(path: str) -> Optional[str]:
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".txt", ".md"):
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        if ext == ".pdf":
            r = PdfReader(path)
            parts = []
            for p in r.pages:
                t = p.extract_text() or ""
                if t.strip():
                    parts.append(t)
            return "\n\n".join(parts)
        if ext == ".docx":
            d = docx.Document(path)
            return "\n".join([p.text for p in d.paragraphs if p.text.strip()])
        if ext == ".xlsx":
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            parts=[]
            for ws in wb.worksheets:
                parts.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    line = "\t".join("" if v is None else str(v) for v in row)
                    if line.strip():
                        parts.append(line)
            return "\n".join(parts)
    except Exception:
        return None
    return None
